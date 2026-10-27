"""M6: LaTeX→OMML, 한글 폰트 4중 고정, DOCX 문서 조립, 출력 후 자동 검사."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any

import docx
import pytest
from lxml import etree
from PIL import Image, ImageDraw

from dawoni.export.docx.builder import (
    DocWriter,
    ExportError,
    FigureContext,
    choose_layout,
    loose_segments,
)
from dawoni.export.docx.check import check_docx
from dawoni.export.docx.documents import (
    KeyOptions,
    PaperOptions,
    build_answer_key,
    build_exam_paper,
    build_grading_report,
    build_score_sheets,
)
from dawoni.export.docx.fonts import FontSpec, add_run, fix_document_fonts
from dawoni.export.fill import FillError, fill_item, format_value
from dawoni.grading import checks
from dawoni.grading.store import ExamStore
from dawoni.mathx.omml import M_NS, W_NS, OmmlConverter, normalize_latex
from dawoni.models.content import MathRun, ParagraphBlock, TextRun, runs_from_text
from dawoni.models.grading import (
    Exam,
    ExamQuestion,
    GradeResult,
    ItemJudgement,
    Rubric,
    RubricDraft,
    ScopedSolution,
    SolutionDraft,
    Transcript,
    TranscriptLine,
)
from dawoni.models.item import ItemContent
from dawoni.workspace import create_workspace

EXPORT_DIR = Path(__file__).parents[1] / "dawoni" / "export"
NS = {"w": W_NS, "m": M_NS}


def T(v: str, **k: Any) -> dict[str, Any]:
    return {"t": "text", "v": v, **k}


def M(v: str) -> dict[str, Any]:
    return {"t": "math", "v": v}


def _xml(path: Path, part: str = "word/document.xml") -> etree._Element:  # pyright: ignore[reportPrivateUsage]
    with zipfile.ZipFile(path) as z:
        return etree.fromstring(z.read(part))


# ---------------------------------------------------------------- 수식


def test_omml_batch_conversion_and_failures() -> None:
    conv = OmmlConverter()
    res = conv.convert_many([r"\frac{1}{2}mv^2", r"v=\sqrt{2gh}\,\text{(속력)}", r"\foo{x}", "x^{"])
    assert res[r"\frac{1}{2}mv^2"].ok
    om = res[r"v=\sqrt{2gh}\,\text{(속력)}"].omml
    assert om is not None and om.find(".//m:rad", NS) is not None, "근호가 OMML 구조로"
    assert not res[r"\foo{x}"].ok and "\\foo" in (res[r"\foo{x}"].error or "")
    assert not res["x^{"].ok
    problems = conv.problems([r"\foo{x}", "x"])
    assert len(problems) == 1 and "표준 LaTeX로 고칠 것" in problems[0]


def test_omml_cache_avoids_second_pandoc_call(monkeypatch: pytest.MonkeyPatch) -> None:
    conv = OmmlConverter()
    calls: list[list[str]] = []
    real = conv._run  # pyright: ignore[reportPrivateUsage]

    def counting(latexes: list[str]) -> Any:
        calls.append(latexes)
        return real(latexes)

    monkeypatch.setattr(conv, "_run", counting)
    conv.convert_many(["a+b", "c"])
    conv.convert_many(["c", "a+b", "d"])
    assert calls == [["a+b", "c"], ["d"]]


@pytest.mark.parametrize(
    ("src", "expected"),
    [
        (r"\ce{H2O}", r"\mathrm{H_{2}O}"),
        (r"\ce{SO4^2-}", r"\mathrm{SO_{4}^{2-}}"),
        (r"\ce{2H2 + O2 -> 2H2O}", r"\mathrm{2H_{2} + O_{2} \rightarrow  2H_{2}O}"),
        (r"30\degree", r"30{}^{\circ}"),
        (r"25\celsius", r"25{}^{\circ}\mathrm{C}"),
    ],
)
def test_normalize_latex(src: str, expected: str) -> None:
    assert normalize_latex(src) == expected
    assert OmmlConverter().convert(src).ok


def test_runs_from_text() -> None:
    runs = runs_from_text(r"속력은 \(v=\sqrt{2gh}\)이고 넓이 $S$는")
    assert [type(r) for r in runs] == [TextRun, MathRun, TextRun, MathRun, TextRun]
    assert runs[1].v == r"v=\sqrt{2gh}" and runs[3].v == "S"
    assert runs_from_text("가격 $5와 $3") == [TextRun(t="text", v="가격 $5와 $3")], (
        "금액은 수식 아님"
    )


@pytest.mark.parametrize(
    ("s", "maths"),
    [
        (r"x좌표는 x=-\frac{b}{2a}=2이다.", [r"x=-\frac{b}{2a}=2"]),
        (r"(단, \frac{1}{2}) 이고 \(y=3\)", [r"\frac{1}{2}", "y=3"]),
        ("식을 y=(x-2)^2+3으로 변형", ["y=(x-2)^2+3"]),
        ("그냥 문장 x=2와 (0,1)", []),
    ],
)
def test_loose_segments(s: str, maths: list[str]) -> None:
    segs = loose_segments(s)
    assert [v for k, v in segs if k == "math"] == maths
    assert "".join(v for _, v in segs).replace(" ", "").startswith(s[:2])


def test_solution_math_validator_requests_repair() -> None:
    d = SolutionDraft.model_validate(
        {
            "steps": [
                {"explanation": r"표준형 \(y=(x-2)^2+3\)", "math": r"\foo{x}", "uses": ["S1"]}
            ],
            "final_answer": {"latex": "3"},
        }
    )
    problems = checks.check_solution_math(d)
    assert len(problems) == 1 and "\\foo" in problems[0]


# ---------------------------------------------------------------- 변수 채우기


def test_fill_item_placeholders(incline_item: dict[str, Any]) -> None:
    incline_item["variables"]["mu"]["decimals"] = 1
    incline_item["stem"][0]["runs"].append(T(r" (마찰 계수 \var{mu}, 무게 \var{m*9.8} N)"))
    item = fill_item(ItemContent.model_validate(incline_item))
    first = item.stem[0]
    assert isinstance(first, ParagraphBlock)
    joined = "".join(r.v for r in first.runs)
    assert r"2\,\mathrm{kg}" in joined and r"30^\circ" in joined
    assert "마찰 계수 0.2, 무게 19.6 N" in joined
    assert "\\var{" not in str(item.model_dump())
    assert item.figures[0].spec is not None and item.constraints, "스펙·제약 식은 그대로"
    assert format_value(3.0) == "3" and format_value(1 / 3) == "0.3333"


def test_fill_item_bad_expression() -> None:
    item = ItemContent.model_validate(
        {
            "item_type": "short_answer",
            "variables": {"a": {"value": 2}},
            "stem": [{"type": "paragraph", "runs": [M(r"\var{a/(a-2)}")]}],
        }
    )
    with pytest.raises(FillError):
        fill_item(item)


# ---------------------------------------------------------------- 폰트 4중 고정 · 검사


def test_only_fonts_module_creates_runs() -> None:
    """글자 런은 fonts.add_run()으로만 (한글 폰트 4중 고정이 빠지는 길을 막는다)."""
    offenders: list[str] = []
    for p in EXPORT_DIR.rglob("*.py"):
        if p.name == "fonts.py":
            continue
        src = p.read_text(encoding="utf-8")
        for m in re.finditer(r"(?<!fonts)\.add_run\(|add_paragraph\(\s*['\"f]", src):
            line = src[: m.start()].count("\n") + 1
            offenders.append(f"{p.name}:{line}")
    assert offenders == []


def test_fixed_fonts_pass_check(tmp_path: Path) -> None:
    w = DocWriter(FontSpec())
    w.prepare_math([r"v=\text{속력}", "x^2"])
    p = w.par()
    w.text(p, "한글 문장 ")
    w.math(p, r"v=\text{속력}")
    w.heading("제목")
    out = w.save(tmp_path / "a.docx")
    rep = check_docx(out)
    assert rep.ok, rep.errors
    assert rep.equations == 1 and rep.hangul_runs >= 2

    styles = _xml(out, "word/styles.xml")
    rf = styles.find("w:docDefaults/w:rPrDefault/w:rPr/w:rFonts", NS)
    assert rf is not None
    assert {k.split("}")[1]: v for k, v in rf.attrib.items()} == {
        "ascii": "맑은 고딕", "hAnsi": "맑은 고딕", "eastAsia": "맑은 고딕", "cs": "맑은 고딕"
    }  # fmt: skip
    assert not [a for el in styles.iter() for a in el.attrib if a.endswith("Theme")
                or a.endswith("cstheme")]  # fmt: skip
    doc = _xml(out)
    mr = [r for r in doc.iter(f"{{{M_NS}}}r") if "속력" in "".join(r.itertext())]
    assert mr and mr[0].find("w:rPr/w:rFonts", NS) is not None
    assert mr[0].find("w:rPr/w:rFonts", NS).get(f"{{{W_NS}}}eastAsia") == "맑은 고딕"  # pyright: ignore[reportOptionalMemberAccess]
    lang = styles.find("w:docDefaults/w:rPrDefault/w:rPr/w:lang", NS)
    assert lang is not None and lang.get(f"{{{W_NS}}}eastAsia") == "ko-KR"


def test_check_catches_tofu_risks(tmp_path: Path) -> None:
    d = docx.Document()  # 고정하지 않은 문서
    p = d.add_paragraph()
    p.add_run("한글")
    p.add_run("$x^2$ 와 \\frac{1}{2}")
    out = tmp_path / "bad.docx"
    d.save(str(out))
    errors = "\n".join(check_docx(out).errors)
    assert "테마 폰트" in errors
    assert "eastAsia 폰트 없음" in errors
    assert "LaTeX 원문" in errors

    # 고정 후에도 런을 직접 만들면 잡힌다
    d2 = docx.Document()
    fix_document_fonts(d2, FontSpec())
    p2 = d2.add_paragraph()
    add_run(p2, "정상", FontSpec())
    p2.add_run("직접 만든 한글 런")
    d2.save(str(out))
    rep = check_docx(out)
    assert len(rep.errors) == 1 and "eastAsia" in rep.errors[0]


# ---------------------------------------------------------------- 선지 배치


def _mc(choices: list[list[dict[str, Any]]], layout: str = "auto") -> ItemContent:
    return ItemContent.model_validate(
        {
            "item_type": "multiple_choice",
            "stem": [{"type": "paragraph", "runs": [T("다음 중 옳은 것은?")]}],
            "choices": [{"label": i, "runs": r} for i, r in enumerate(choices, start=1)],
            "choice_layout": layout,
        }
    )


def test_choose_layout() -> None:
    col = 8.6
    short = _mc([[M(str(v))] for v in (2, 4, 6, 8, 10)])
    assert choose_layout(short, col, 10, []) == ("row", 5)
    mid = _mc([[T(s)] for s in ("ㄱ", "ㄴ", "ㄱ, ㄷ", "ㄴ, ㄷ", "ㄱ, ㄴ, ㄷ")])
    assert choose_layout(mid, col, 10, []) == ("grid", 3)
    long = _mc([[T("물체에 작용하는 알짜힘은 항상 0이다")] for _ in range(5)])
    assert choose_layout(long, col, 10, []) == ("column", 1)
    assert choose_layout(_mc([[T("ㄱ")]] * 5, "column"), col, 10, []) == ("column", 1)
    assert choose_layout(_mc([[T("ㄱ")]] * 5, "grid_2"), col, 10, []) == ("grid", 2)
    assert choose_layout(short, col, 10, [2.5] * 5)[1] == 3, "그래프 선지 2.5cm → 3+2"
    assert choose_layout(short, col, 10, [1.3] * 5) == ("row", 5)
    assert choose_layout(short, col, 10, [6.0] * 5) == ("column", 1)


# ---------------------------------------------------------------- 시험지


def _graph(path: Path, k: int, px: int = 300) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    im = Image.new("RGB", (px, px * 4 // 5), "white")
    d = ImageDraw.Draw(im)
    d.line([(10, px * 3 // 4), (px - 10, px * 3 // 4)], fill="black", width=3)
    d.line([(20, 10), (20 + k * 30, px * 3 // 4)], fill="black", width=3)
    im.save(path, dpi=(300, 300))


@pytest.fixture
def sample_exam(tmp_path: Path) -> tuple[Exam, Path]:
    exam_dir = tmp_path / "exam"
    for k in range(1, 6):
        _graph(exam_dir / "figures" / "3" / f"F{k}.png", k)
    _graph(exam_dir / "figures" / "4" / "F1.png", 1, px=900)
    qs: list[dict[str, Any]] = [
        {"number": "1", "points": 2, "item": {
            "item_type": "multiple_choice",
            "stem": [{"type": "paragraph", "runs": [M(r"\left(\frac{1}{2}\right)^{-2}"), T("의 값은?")]}],
            "choices": [{"label": i, "runs": [M(str(v))]} for i, v in enumerate([2, 4, 6, 8, 10], 1)],
            "answer": {"kind": "choice", "labels": [2]}}},
        {"number": "2", "points": 3, "item": {
            "item_type": "multiple_choice",
            "stem": [
                {"type": "paragraph", "runs": [T("옳은 것", bold=True, underline=True), T("만을 <보기>에서 고른 것은?")]},
                {"type": "bogi", "items": [
                    {"label": "ㄱ", "runs": [T("가속도는 "), M(r"g\sin\theta"), T("이다.")]},
                    {"label": "ㄴ", "runs": [T("속력 "), M(r"v=\sqrt{2gh}\,\text{(속력)}")]}]}],
            "choices": [{"label": i, "runs": [T(s)]} for i, s in enumerate(["ㄱ", "ㄴ", "ㄱ, ㄴ", "없음", "모두"], 1)]}},
        {"number": "3", "points": 4, "item": {
            "item_type": "multiple_choice",
            "stem": [{"type": "paragraph", "runs": [M("y=2^{x}"), T("의 그래프는?")]}],
            "choices": [{"label": i, "figure_id": f"F{i}", "note": f"그래프 {i}"} for i in range(1, 6)],
            "figures": [{"id": f"F{i}", "mode": "raster", "orig": f"figures/3/F{i}.png"} for i in range(1, 6)]}},
        {"number": "4", "points": 3, "item": {
            "item_type": "multiple_choice",
            "stem": [{"type": "paragraph", "runs": [T("옳은 그래프는?")]}],
            "choices": [{"label": i, "note": f"그래프 {i}"} for i in range(1, 6)],
            "choices_figure_id": "F1",
            "figures": [{"id": "F1", "mode": "raster", "orig": "figures/4/F1.png"}]}},
        {"number": "5", "points": 3, "item": {
            "item_type": "short_answer",
            "stem": [
                {"type": "paragraph", "runs": [T("물질 "), M(r"\ce{H2O}"), T("의 자료이다.")]},
                {"type": "table", "header_rows": 1, "rows": [[[T("물질")], [T("상태")]], [[M(r"\ce{H2O}")], [T("액체")]]]},
                {"type": "box", "title": "자료", "paragraphs": [[M(r"\ce{2H2 + O2 -> 2H2O}")]]},
                {"type": "conditions", "items": [{"label": "(가)", "runs": [M("f(1)=3")]}]},
                {"type": "display_math", "latex": r"\int_0^1 f(x)\,dx"}]}},
        {"number": "서술형1", "points": 6, "item": {
            "item_type": "essay",
            "stem": [{"type": "paragraph", "runs": [M(r"0 \le x \le 3"), T("에서 최댓값을 구하시오.")]}]}},
    ]  # fmt: skip
    exam = Exam(
        title="2학기 중간고사",
        scope_id="sc",
        questions=[ExamQuestion.model_validate(q) for q in qs],
    )
    return exam, exam_dir


def test_exam_paper(sample_exam: tuple[Exam, Path], tmp_path: Path) -> None:
    exam, exam_dir = sample_exam
    w = build_exam_paper(exam, exam_dir, FontSpec(), PaperOptions(subtitle="고1"))
    out = w.save(tmp_path / "paper.docx")
    rep = check_docx(out)
    assert rep.ok, rep.errors
    assert w.warnings == []
    assert rep.images == 6, "그림 선지 5 + 한 장짜리 선지 그림 1"
    assert rep.equations >= 13

    doc = _xml(out)
    text = "".join(doc.itertext())
    assert "< 보 기 >" in text and "[3점]" in text and "[서술형1]" in text and "(풀이)" in text
    # 배점은 <보기> 상자 뒤가 아니라 발문 끝에
    first_par = next(p for p in doc.iter(f"{{{W_NS}}}p") if "<보기>에서" in "".join(p.itertext()))
    assert "".join(first_par.itertext()).endswith("[3점]")
    cols = [c for c in doc.iter(f"{{{W_NS}}}cols") if c.get(f"{{{W_NS}}}num") == "2"]
    assert cols and cols[0].get(f"{{{W_NS}}}sep") == "1", "2단 + 가운데 선"
    # 그래프 선지: 3+2 배치 표 (테두리 없음)
    grids = [len(t.findall("w:tblGrid/w:gridCol", NS)) for t in doc.iter(f"{{{W_NS}}}tbl")]
    assert 3 in grids and 5 in grids, "숫자 선지는 한 줄 5개, 그래프 선지·<보기> 선지는 3열"
    # 한 장짜리 선지 그림 문항은 ①~⑤를 따로 쓰지 않는다
    q4 = text.split("옳은 그래프는?")[1].split("물질")[0]
    assert "①" not in q4
    # 자동 간격 끔 ("[3점]"이 "[3 점]"으로 보이지 않게)
    styles = _xml(out, "word/styles.xml")
    ad = styles.find("w:docDefaults/w:pPrDefault/w:pPr/w:autoSpaceDN", NS)
    assert ad is not None and ad.get(f"{{{W_NS}}}val") == "0"


def test_exam_paper_refuses_unconvertible_math(
    sample_exam: tuple[Exam, Path], tmp_path: Path
) -> None:
    exam, exam_dir = sample_exam
    bad = exam.questions[0].model_copy(deep=True)
    first = bad.item.stem[0]
    assert isinstance(first, ParagraphBlock)
    first.runs[0] = MathRun(t="math", v=r"\foo{x}")
    exam2 = exam.model_copy(update={"questions": [bad]})
    with pytest.raises(ExportError, match="foo"):
        build_exam_paper(exam2, exam_dir, FontSpec())


def test_missing_figure_is_warned(tmp_path: Path) -> None:
    item = ItemContent.model_validate(
        {
            "item_type": "short_answer",
            "stem": [{"type": "figure", "figure_id": "F1"}],
            "figures": [{"id": "F1", "mode": "raster", "orig": "figures/none.png"}],
        }
    )
    w = DocWriter()
    w.question("1", 3, item, FigureContext.of(item, tmp_path))
    assert w.warnings and "원본 이미지가 없어" in w.warnings[0]
    assert check_docx(w.save(tmp_path / "x.docx")).ok


# ---------------------------------------------------------------- 해설지 · 리포트 · 성적표


@pytest.fixture
def graded_store(tmp_path: Path, sample_exam: tuple[Exam, Path]) -> tuple[Exam, ExamStore]:
    ws = create_workspace(tmp_path / "ws")
    exam = Exam(
        title="채점 시험",
        scope_id="sc",
        questions=[sample_exam[0].questions[-1]],  # 서술형1
    )
    store = ExamStore.create(ws, exam)
    store.save_solution(
        ScopedSolution(
            **SolutionDraft.model_validate(
                {
                    "steps": [
                        {
                            "explanation": r"완전제곱식 \(y=-(x-1)^2+4\)",
                            "math": "y=-(x-1)^2+4",
                            "uses": ["S1"],
                        },
                        {"explanation": "꼭짓점 x=1이 범위 안", "math": "x=1", "uses": ["C3"]},
                    ],
                    "final_answer": {"latex": r"\text{최댓값 } 4"},
                    "alternatives": [
                        {
                            "summary": r"꼭짓점 공식 x=-\frac{b}{2a}",
                            "steps": [{"explanation": "대입", "uses": ["S2"]}],
                        }
                    ],
                }
            ).model_dump(),
            question_number="서술형1",
            scope_version=1,
        )
    )
    store.save_rubric(
        Rubric(
            **RubricDraft.model_validate(
                {
                    "items": [
                        {
                            "id": "R1",
                            "points": 3,
                            "criterion": "y=-(x-1)^2+4로 변형",
                            "evidence_hint": r"-\left(x-1\right)^2+4",
                        },
                        {"id": "R2", "points": 3, "criterion": "최댓값 4", "evidence_hint": "4"},
                    ],
                    "never_deduct": ["미분을 쓰지 않았다는 이유"],
                }
            ).model_dump(),
            question_number="서술형1",
            status="approved",
        )
    )
    for st, total, status in (("A", 6, "auto"), ("B", 3, "needs_review")):
        store.save_transcript(
            Transcript(
                student_id=st,
                question_number="서술형1",
                lines=[
                    TranscriptLine(math="y=-(x-1)^2+4"),
                    TranscriptLine(text=r"따라서 \(x=1\)일 때 최댓값 4", struck=st == "B"),
                    TranscriptLine(math=r"\frac{1}{", uncertain=["분모"]),
                ],
                confirmed=True,
            )
        )
        store.save_grade(
            GradeResult(
                student_id=st,
                question_number="서술형1",
                rubric_version=1,
                judgements=[
                    ItemJudgement(
                        rubric_item_id="R1",
                        awarded=3,
                        evidence_quote="y=-(x-1)^2+4",
                        reason="표준형 변형",
                        confidence=0.9,
                    ),
                    ItemJudgement(
                        rubric_item_id="R2",
                        awarded=total - 3,
                        reason="결론",
                        confidence=0.9,
                        deduction_reason=None if total == 6 else "최댓값을 지움",
                    ),
                ],
                comment="잘 했습니다." if total == 6 else "결론을 지우지 마세요.",
                flags=[] if total == 6 else ["판독 불확실"],
                total=total,
                max_points=6,
                status=status,  # pyright: ignore[reportArgumentType]
            )
        )
    return exam, store


def test_answer_key(graded_store: tuple[Exam, ExamStore], tmp_path: Path) -> None:
    exam, store = graded_store
    w = build_answer_key(exam, store, FontSpec(), KeyOptions(with_questions=True, with_rubric=True))
    out = w.save(tmp_path / "key.docx")
    rep = check_docx(out)
    assert rep.ok, rep.errors
    assert rep.latex_fallbacks == 0, "AI 설명 속 구분자 없는 식도 수식으로"
    text = "".join(_xml(out).itertext())
    for s in ("정답 및 해설", "문항 서술형1", "따라서 답은", "[다른 풀이]", "채점 기준 (v1, 승인됨)",
              "감점하지 않는 경우", "미분을 쓰지 않았다는 이유"):  # fmt: skip
        assert s in text, s


def test_grading_report_and_score_sheets(
    graded_store: tuple[Exam, ExamStore], tmp_path: Path
) -> None:
    exam, store = graded_store
    w = build_grading_report(exam, store, FontSpec())
    out = w.save(tmp_path / "report.docx")
    rep = check_docx(out)
    assert rep.ok, rep.errors
    assert w.math_fallbacks == [r"\frac{1}{"] * 2, "학생이 쓴 깨진 식은 원문 그대로 (두 학생)"
    assert rep.latex_fallbacks == 2
    doc = _xml(out)
    text = "".join(doc.itertext())
    assert "학생 A" in text and "학생 B" in text and "3 △" in text and "※ 판독 불확실" in text
    assert "(불확실: 분모)" in text
    struck = [r for r in doc.iter(f"{{{W_NS}}}r") if r.find("w:rPr/w:strike", NS) is not None]
    assert struck, "지운 줄은 취소선"
    assert not re.search("[☀-➿]", text), "맑은 고딕에 없는 기호(⚠ 등) 금지"

    w2 = build_score_sheets(exam, store, FontSpec())
    out2 = w2.save(tmp_path / "scores.docx")
    assert check_docx(out2).ok
    text2 = "".join(_xml(out2).itertext())
    assert "성적표" in text2 and "학생: B" in text2 and "최댓값을 지움" in text2
    assert "판독 불확실" not in text2 and "△" not in text2, "학생용에는 내부 표시 없음"
