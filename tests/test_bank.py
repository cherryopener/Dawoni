"""문제은행: 공통 단원표, 문항 저장(팀 공유 충돌), 구간 DB화 + 해설 채우기·원본 비교."""

from __future__ import annotations

import json
from collections.abc import Callable
from datetime import timedelta
from pathlib import Path
from typing import Any

import pymupdf
import pytest
from conftest import MakeFake
from tests_support import fake_prompts

from dawoni.bank import BankError
from dawoni.bank import ingest as bi
from dawoni.bank.store import BankStore
from dawoni.bank.units import UnitStore, merge, parse_toc
from dawoni.config import AppConfig, save_config
from dawoni.fsio import read_json
from dawoni.grading.store import ExamStore, ScopeStore
from dawoni.ingest.answer_key import KeyIngest
from dawoni.ingest.exam import IngestResult
from dawoni.models.answer_key import AnswerKey, KeyPart
from dawoni.models.bank import AnswerComparison, BankItem, BankSourceRef, OriginalSolution
from dawoni.models.base import utcnow
from dawoni.models.grading import ExamQuestion, ScopedSolution
from dawoni.models.item import ChoiceAnswer
from dawoni.models.scope import Concept, ForbiddenTechnique, ScopeProfile, Skill
from dawoni.models.source import SourceChunk
from dawoni.paths import AppPaths
from dawoni.services import Services, open_services
from dawoni.workspace import create_workspace

# ---------------------------------------------------------------- 단원표


def test_default_trees_seeded_once(tmp_path: Path) -> None:
    ws = create_workspace(tmp_path / "ws")
    store = UnitStore(ws)
    trees = store.all()
    assert {t.id for t in trees} >= {"phy1", "phy2", "chem1", "bio1", "earth1"}
    assert not any(t.reviewed for t in trees), "기본 초안은 검토 필요"
    phy1 = store.get("phy1")
    assert phy1.label("phy1.1.1.2") == "물리학 I > 역학과 에너지 > 힘과 운동 > 뉴턴 운동 법칙"
    assert phy1.descendants("phy1.1.1") == {"phy1.1.1", "phy1.1.1.1", "phy1.1.1.2", "phy1.1.1.3"}
    (ws.units / "phy2.json").unlink()
    assert "phy2" not in {t.id for t in store.all()}, "강사가 지운 기본 단원표는 되살리지 않음"


def test_parse_toc() -> None:
    toc = """Ⅰ. 역학과 에너지 ........ 8
      1. 힘과 운동  10
        (1) 물체의 운동 12
        (2) 뉴턴 운동 법칙 … 20
      2. 에너지와 열 40
    Ⅱ. 전기와 자기
      1. 전기장"""
    units = parse_toc(toc, "gift")
    assert [(u.id, u.title) for u in units] == [
        ("gift.1", "역학과 에너지"),
        ("gift.2", "전기와 자기"),
    ]
    force = units[0].children[0]
    assert (force.id, force.title) == ("gift.1.1", "힘과 운동")
    assert [c.title for c in force.children] == ["물체의 운동", "뉴턴 운동 법칙"]
    with pytest.raises(BankError):
        parse_toc("\n\n", "x")


def test_merge_keeps_old_ids_resolvable(tmp_path: Path) -> None:
    store = UnitStore(create_workspace(tmp_path / "ws"))
    t = store.get("phy1")
    m = merge(t, "phy1.1.2", "phy1.1.1")
    assert m.resolve("phy1.1.2") == "phy1.1.1" and m.resolve("phy1.1.2.1") == "phy1.1.2.1"
    assert "phy1.1.2.1" in m.descendants("phy1.1.1"), "아래 단원은 합친 단원 밑으로"
    with pytest.raises(BankError, match="자기 아래"):
        merge(t, "phy1.1", "phy1.1.1")
    saved = store.save(m, base_version=t.version, by="김선생", in_use={"phy1.1.2"})
    assert saved.version == t.version + 1 and saved.updated_by == "김선생"


def test_save_refuses_stale_and_used_removal(tmp_path: Path) -> None:
    store = UnitStore(create_workspace(tmp_path / "ws"))
    t = store.get("phy1")
    store.save(t.model_copy(update={"subject": "물리학1"}), base_version=t.version)
    with pytest.raises(BankError, match="먼저"):
        store.save(t, base_version=t.version)
    t = store.get("phy1")
    pruned = t.model_copy(update={"units": t.units[1:]})
    with pytest.raises(BankError, match="문항이 들어 있는 단원"):
        store.save(pruned, base_version=t.version, in_use={"phy1.1.1.2"})
    assert store.save(pruned, base_version=t.version).version == t.version + 1


# ---------------------------------------------------------------- 원본 해설 · 비교 (AI 없음)


def _kp(
    answer: str | None, solution: str = "", expr: str | None = None, label: str = ""
) -> KeyPart:
    return KeyPart(
        label=label, answer=answer, answer_expr=expr, solution=solution, credit_policy="unspecified"
    )


def test_original_level() -> None:
    assert bi.original_level([]) == "none"
    assert bi.original_level([_kp(None)]) == "none"
    assert bi.original_level([_kp("③")]) == "answer_only"
    assert bi.original_level([_kp("2", "F=ma 이용")]) == "brief"
    assert bi.original_level([_kp("2", "가" * bi.BRIEF_CHARS)]) == "full"


def test_choice_numbers() -> None:
    assert bi.choice_numbers("③") == {3}
    assert bi.choice_numbers("\\text{③}") == {3}
    assert bi.choice_numbers("정답 3번") == {3}
    assert bi.choice_numbers("①, ④") == {1, 4}
    assert bi.choice_numbers("2.5") == set()


# ---------------------------------------------------------------- 구간 DB화


def _mc(number: str) -> ExamQuestion:
    return ExamQuestion.model_validate(
        {
            "number": number,
            "points": 2,
            "item": {
                "item_type": "multiple_choice",
                "number_label": number,
                "stem": [
                    {"type": "paragraph", "runs": [{"t": "text", "v": "운동량이 보존되는 경우는?"}]}
                ],
                "choices": [
                    {"label": i, "runs": [{"t": "text", "v": f"보기{i}"}]} for i in range(1, 6)
                ],
            },
        }
    )


def _short(number: str) -> ExamQuestion:
    return ExamQuestion.model_validate(
        {
            "number": number,
            "points": 3,
            "item": {
                "item_type": "short_answer",
                "number_label": number,
                "stem": [
                    {
                        "type": "paragraph",
                        "runs": [{"t": "text", "v": "물체의 가속도의 크기를 구하시오."}],
                    }
                ],
            },
        }
    )


def _sol(answer: str, uses: list[str], expr: str | None = None) -> dict[str, Any]:
    return {
        "steps": [{"explanation": "운동 방정식을 세운다", "math": "F=ma", "uses": uses}],
        "final_answer": {"latex": answer, **({"expr": expr} if expr else {})},
    }


def _pdf(path: Path) -> Path:
    doc = pymupdf.open()
    p = doc.new_page()
    for y, s in [
        (100, "1. 운동량이 보존되는 경우는?"),
        (300, "2. 물체의 가속도의 크기를 구하시오."),
    ]:
        p.insert_text((50, y), s, fontname="korea", fontsize=10)  # pyright: ignore[reportUnknownMemberType]
    k = doc.new_page()
    k.insert_text((50, 100), "정답 및 해설  1. ③  2. 2", fontname="korea", fontsize=10)  # pyright: ignore[reportUnknownMemberType]
    doc.save(str(path))  # pyright: ignore[reportUnknownMemberType]
    return path


Setup = Callable[[list[dict[str, Any]]], tuple[Services, str, Path]]


@pytest.fixture
def setup(
    app_paths: AppPaths, tmp_path: Path, make_fake: MakeFake, monkeypatch: pytest.MonkeyPatch
) -> Setup:
    """문항 구조화·해설지 읽기는 시험지 DB화 테스트가 검사하므로 여기서는 결과를 넣어 둔다."""

    def fake_extract(*_: Any, **__: Any) -> IngestResult:
        return IngestResult(questions=[_mc("1"), _short("2")], figure_files={})

    def fake_key(*_: Any, **__: Any) -> KeyIngest:
        return KeyIngest(
            keys=[
                AnswerKey(question_number="1", source="t", parts=[_kp("③")]),
                AnswerKey(question_number="2", source="t", parts=[_kp("2", "F=ma", expr="2")]),
            ]
        )

    monkeypatch.setattr("dawoni.ingest.exam.extract_exam", fake_extract)
    monkeypatch.setattr("dawoni.ingest.answer_key.extract_answer_key", fake_key)

    def make(script: list[dict[str, Any]]) -> tuple[Services, str, Path]:
        exe = make_fake(
            "script", script=[{"response": json.dumps(r, ensure_ascii=False)} for r in script]
        )
        cfg = AppConfig()
        cfg.agents["agy"].executable = str(exe)
        cfg.agents["claude"].enabled = False
        cfg.agents["codex"].enabled = False
        cfg.parallel_questions = 1  # 가짜 AI 응답 순서를 고정
        save_config(app_paths, cfg)
        ws = create_workspace(tmp_path / "ws")
        svc = open_services(ws.root, paths=app_paths)
        src = bi.add_source(
            svc,
            str(_pdf(tmp_path / "문제집.pdf")),
            "problem_bank",
            title="OO 물리 문제집",
            publisher="OO학원",
            unit_trees=["phy1"],
        )
        bi.add_chunks(
            svc, src, [SourceChunk(index=1, pages=[1], key_pages=[2], unit_hint="phy1.1.1")]
        )
        return svc, src.source_id, exe

    return make


def test_ingest_chunk_compares_with_original(setup: Setup) -> None:
    svc, sid, exe = setup(
        [
            _sol("③", ["C5"]),  # 1번: 원본 답 ③과 같음
            _sol("4", ["C4", "C4", "C3"], "4"),  # 2번: 원본 2 ≠ AI 4 → 둘 다 검산
            {"verdict": "valid", "reasoning": "a=2가 조건을 만족"},
            {"verdict": "invalid", "reasoning": "질량을 두 번 나눔"},
        ]
    )
    res = bi.ingest_chunk(svc, sid, 1, by="김선생", machine="pc-a", solve=True)
    assert len(fake_prompts(exe)) == 4, "원본 답이 있으면 한 번만 풀고, 다를 때만 검산"
    one, two = res.items
    assert one.answer_check and one.answer_check.verdict == "same" and not one.needs_check
    assert one.original and one.original.level == "answer_only" and one.solution_by == "ai"
    assert one.unit_ids == ["phy1.1.1", "phy1.1.1.3"], "구간 단원 힌트 + 풀이에 쓴 단원"
    assert two.needs_check and two.answer_check and two.answer_check.verdict == "differs"
    assert "원본 2 (검산 valid" in two.answer_check.notes[0]
    assert two.original and two.original.level == "brief"
    assert two.source.display() == "OO 물리 문제집 1쪽 2번" and two.created_by == "김선생"

    src = bi.get_source(svc.workspace, sid)
    assert src.chunks[0].status == "ingested" and src.ingest.item_count == 2
    assert src.chunks[0].machine is None

    again = bi.ingest_chunk(svc, sid, 1, machine="pc-a")
    assert again.items == [] and len(again.skipped) == 2, "이어 하기: 있는 문항은 덮어쓰지 않음"
    assert len(fake_prompts(exe)) == 4


def test_no_original_answer_solves_twice(setup: Setup) -> None:
    svc, sid, exe = setup(
        [_sol("③", ["C5"]), _sol("②", ["C5"]), _sol("2", ["C4"], "2"), _sol("2", ["C4"], "2")]
    )
    src = bi.get_source(svc.workspace, sid)
    src = bi.add_chunks(svc, src, [SourceChunk(index=1, pages=[1])])  # 해설 쪽 없음
    res = bi.ingest_chunk(svc, sid, 2, solve=True)
    one, two = res.items
    assert one.needs_check and one.answer_check and one.answer_check.verdict == "unsure"
    assert two.answer_check and two.answer_check.verdict == "no_original" and not two.needs_check
    assert two.original and two.original.level == "none"
    assert len(fake_prompts(exe)) == 4


def test_ingest_without_solving_then_solve_on_request(setup: Setup) -> None:
    from dawoni.sharing import WorkInProgress, work_mark

    svc, sid, exe = setup(
        [
            _sol("③", ["C5"]),
            _sol("4", ["C4", "C4", "C3"], "4"),
            {"verdict": "valid", "reasoning": "a=2가 조건을 만족"},
            {"verdict": "invalid", "reasoning": "질량을 두 번 나눔"},
        ]
    )
    res = bi.ingest_chunk(svc, sid, 1, by="김선생", machine="pc-a")
    assert fake_prompts(exe) == [], "DB화만: AI 풀이 없음 (문항 구조화·해설지 읽기만)"
    one, two = res.items
    assert one.solution is None and one.answer_check is None and not one.needs_check
    assert one.unit_ids == ["phy1.1.1"], "단원은 구간 단원"
    assert one.original and one.original.parts[0].answer == "③", "원본 답은 옮겨 둔다"

    # 다른 PC가 풀고 있는 문항은 막는다
    bank = BankStore(svc.workspace)
    with (
        work_mark(bank.item_dir(one.id), "solve", "solve", "pc-b", "이선생"),
        pytest.raises(WorkInProgress, match="이선생 선생님이 지금 만들고"),
    ):
        bi.solve_items(svc, [one.id], machine="pc-a")
    assert fake_prompts(exe) == []

    out = bi.solve_items(svc, [one.id, two.id], by="이선생", machine="pc-a")
    assert len(fake_prompts(exe)) == 4 and out.skipped == []
    one, two = out.solved
    assert one.answer_check and one.answer_check.verdict == "same"
    assert one.unit_ids == ["phy1.1.1", "phy1.1.1.3"], "풀고 나면 풀이에 쓴 단원도"
    assert two.needs_check and two.updated_by == "이선생" and two.created_by == "김선생"
    assert bank.get(two.id).solution is not None, "저장됨"

    again = bi.solve_items(svc, [one.id, two.id], machine="pc-a")
    assert again.solved == [] and again.skipped == [one.id, two.id], "이미 풀었으면 건너뜀"
    assert len(fake_prompts(exe)) == 4


def test_other_pc_lock(setup: Setup) -> None:
    svc, sid, _ = setup([])
    src = bi.get_source(svc.workspace, sid)
    busy = src.chunks[0].model_copy(
        update={"status": "ingesting", "machine": "pc-b", "by": "이선생", "updated_at": utcnow()}
    )
    bi.update_source(svc.workspace, src.model_copy(update={"chunks": [busy]}))
    with pytest.raises(BankError, match="이선생 선생님이 이 자료를 DB화"):
        bi.ingest_chunk(svc, sid, 1, machine="pc-a")
    stale = busy.model_copy(update={"updated_at": utcnow() - bi.LOCK_FRESH - timedelta(minutes=1)})
    bi.update_source(svc.workspace, src.model_copy(update={"chunks": [stale]}))
    bi._check_lock(bi.get_source(svc.workspace, sid), "pc-a")  # pyright: ignore[reportPrivateUsage]


def test_bad_chunk_refused(setup: Setup) -> None:
    svc, sid, _ = setup([])
    src = bi.get_source(svc.workspace, sid)
    with pytest.raises(BankError, match="없는 쪽"):
        bi.add_chunks(svc, src, [SourceChunk(index=1, pages=[9])])
    with pytest.raises(BankError, match="없는 단원"):
        bi.add_chunks(svc, src, [SourceChunk(index=1, pages=[1], unit_hint="phy1.9")])


# ---------------------------------------------------------------- 저장 (팀 공유)


def _item(item_id: str) -> BankItem:
    return BankItem(
        id=item_id,
        question=_mc("1"),
        unit_ids=["phy1.1.1"],
        source=BankSourceRef(title="기출", kind="mock_exam", pages=[3], number_label="1"),
    )


def test_store_revision_conflict_and_cache(tmp_path: Path, app_paths: AppPaths) -> None:
    ws = create_workspace(tmp_path / "ws")
    store = BankStore(ws, app_paths.cache)
    store.save(_item("bk-a1"), by="김선생")
    first = store.get("bk-a1")
    edited = store.save(
        first.model_copy(update={"difficulty": 3}), by="이선생", base_revision=first.revision
    )
    assert edited.revision == 2 and edited.created_by == "김선생" and edited.updated_by == "이선생"
    with pytest.raises(BankError, match="이선생"):
        store.save(first.model_copy(update={"difficulty": 4}), base_revision=first.revision)
    rows = store.summaries()
    assert [(r.id, r.source, r.difficulty) for r in rows] == [("bk-a1", "기출 3쪽 1번", 3)]
    assert store.summaries() == rows and list(app_paths.cache.glob("bank_*.json")), "PC별 캐시"
    assert not list(ws.bank.glob("*.json")), "공유 폴더에 공용 목록 파일을 두지 않음"
    assert store.used_unit_ids() == {"phy1.1.1"}


def test_solution_scope_ids_map_to_units() -> None:
    from dawoni.bank.curriculum import default_trees

    scope, mapping = bi.unit_scope([t for t in default_trees() if t.id == "phy1"])
    sol = ScopedSolution.model_validate(
        {**_sol("x", ["C4", "C3", "C4"]), "question_number": "1", "scope_version": 1}
    )
    assert bi.units_from_solution(sol, mapping) == [mapping["C4"], mapping["C3"]]
    assert scope.level.startswith("고등학교 물리학 I")


# ---------------------------------------------------------------- DOCX (출처 없음)


def test_export_without_source(tmp_path: Path, app_paths: AppPaths) -> None:
    from docx import Document

    from dawoni.bank.export import bank_exam, export_bank
    from dawoni.export.docx.fonts import FontSpec

    ws = create_workspace(tmp_path / "ws")
    store = BankStore(ws, app_paths.cache)
    a = _item("bk-a1").model_copy(
        update={
            "original": OriginalSolution(level="answer_only", parts=[_kp("④")]),
            "solution": ScopedSolution.model_validate(
                {**_sol("③", ["C5"]), "question_number": "17", "scope_version": 1}
            ),
            "answer_check": AnswerComparison(verdict="differs", original_answer="④", ai_answer="③"),
        }
    )
    b = _item("bk-b2").model_copy(
        update={
            "question": _short("9"),
            "original": OriginalSolution(level="full", parts=[_kp("2", "뉴턴 제2법칙 " * 10, "2")]),
            "solution": ScopedSolution.model_validate(
                {**_sol("2", ["C4"], "2"), "question_number": "9", "scope_version": 1}
            ),
            "solution_by": "source",
            "answer_check": AnswerComparison(verdict="same"),
        }
    )
    for it in (a, b):
        store.save(it)

    exam, sols, warnings = bank_exam([a, b], "단원 평가", "bank", ws.bank, tmp_path / "out")
    assert [q.number for q in exam.questions] == ["1", "2"]
    assert exam.questions[0].item.answer == ChoiceAnswer(kind="choice", labels=[4]), "원본 답 우선"
    assert "1" not in sols, "원본 답과 다른 AI 해설은 넣지 않음"
    assert "2" not in sols and exam.questions[1].item.solution, "충분한 원본 해설은 원본으로"
    assert warnings and "확인이 필요한" in warnings[0]

    resolved = a.model_copy(
        update={
            "answer_check": a.answer_check.model_copy(  # type: ignore[union-attr]
                update={"resolved": "ai", "resolved_by": "김선생"}
            )
        }
    )
    assert not resolved.needs_check
    exam, sols, _ = bank_exam([resolved], "t", "bank", ws.bank, tmp_path / "out2")
    assert exam.questions[0].item.answer == ChoiceAnswer(kind="choice", labels=[3]) and "1" in sols

    for kind in ("paper", "key"):
        out = export_bank(ws, [a, b], kind, "단원 평가", FontSpec(), tmp_dir=app_paths.cache)
        text = "\n".join(p.text for p in Document(str(out.path)).paragraphs)
        tables = " ".join(
            c.text for t in Document(str(out.path)).tables for r in t.rows for c in r.cells
        )
        assert "기출" not in text + tables and "17" not in text, f"{kind}: 출처·원래 번호 없음"


# ---------------------------------------------------------------- API


def test_api_units_items_export(app_paths: AppPaths, tmp_path: Path) -> None:
    from test_server import _client  # pyright: ignore[reportPrivateUsage]

    c = _client(app_paths)
    c.post("/api/workspace", json={"path": str(tmp_path / "ws"), "create": True})
    trees = c.get("/api/units").json()
    phy1 = next(t for t in trees if t["id"] == "phy1")

    toc = "1. 힘과 운동\n  (1) 등속 원운동\n2. 전기장"
    r = c.post(
        "/api/units", json={"id": "gift", "subject": "영재 물리", "group": "물리", "toc": toc}
    )
    assert r.status_code == 200 and r.json()["units"][0]["children"][0]["id"] == "gift.1.1"
    assert (
        c.post("/api/units", json={"id": "gift", "subject": "x", "group": "x"}).status_code == 400
    )

    ws = create_workspace(tmp_path / "ws")
    BankStore(ws).save(
        _item("bk-a1").model_copy(
            update={
                "original": OriginalSolution(level="answer_only", parts=[_kp("④")]),
                "answer_check": AnswerComparison(
                    verdict="differs", original_answer="④", ai_answer="③"
                ),
            }
        ),
        by="김선생",
    )

    # 문항이 있는 단원은 지울 수 없고, 합치기는 된다
    pruned = {**phy1, "units": phy1["units"][1:]}
    r = c.put("/api/units/phy1", json={"tree": pruned, "base_version": phy1["version"]})
    assert r.status_code == 400 and "합치세요" in r.json()["detail"]
    r = c.post(
        "/api/units/phy1/merge",
        json={
            "source": "phy1.1.1",
            "into": "phy1.1.2",
            "base_version": phy1["version"],
            "by": "이선생",
        },
    )
    assert r.status_code == 200 and r.json()["aliases"] == {"phy1.1.1": "phy1.1.2"}
    r = c.post(
        "/api/units/phy1/merge",
        json={"source": "phy1.1.2", "into": "phy1.1.3", "base_version": phy1["version"]},
    )
    assert r.status_code == 400 and "이선생" in r.json()["detail"], "옛 판으로 고치면 거부"

    rows = c.get("/api/bank/items").json()
    assert [(x["id"], x["needs_check"], x["solved"]) for x in rows] == [("bk-a1", True, False)]
    r = c.post("/api/bank/items/solve", json={"ids": []})
    assert r.status_code == 400 and "풀 문항" in r.json()["detail"]
    assert c.post("/api/bank/sources/none/chunks/1/solve", json={}).status_code == 404
    got = c.get("/api/bank/items/bk-a1").json()
    assert got["unit_labels"]["phy1.1.1"].endswith("에너지와 열"), "합친 단원 이름으로 보임"
    r = c.post(
        "/api/bank/items/bk-a1/resolve",
        json={"side": "original", "base_revision": 1, "by": "김선생"},
    )
    assert r.status_code == 200 and r.json()["answer_check"]["resolved"] == "original"
    assert not c.get("/api/bank/items").json()[0]["needs_check"]

    out = c.post("/api/bank/export", json={"ids": ["bk-a1"], "kind": "paper", "title": "단원 평가"})
    assert out.status_code == 200 and out.json()["path"].startswith("exports/")

    scope = read_json(Path(__file__).parent / "fixtures" / "grading" / "scope.json")
    ScopeStore(ws).save(ScopeProfile.model_validate(scope))
    r = c.post(
        "/api/bank/to-exam", json={"ids": ["bk-a1"], "title": "평가", "scope_id": scope["id"]}
    )
    store = ExamStore(ws, r.json()["exam_id"])
    assert store.exam().questions[0].number == "1"
    key = store.answer_key("1")
    assert key is not None and key.parts[0].answer == "④", "원본 답은 해설지로"


# ---------------------------------------------------------------- 단원 체크 → 범위


def _scope(title: str, units: list[str], concepts: list[str], **kw: Any) -> ScopeProfile:
    return ScopeProfile(
        title=title,
        level="고등학교 물리학 I",
        unit_ids=units,
        concepts=[
            Concept(id=f"C{i}", name=n, summary=f"{n} 설명") for i, n in enumerate(concepts, 1)
        ],
        **kw,
    )


def test_compose_scope_across_textbooks() -> None:
    from dawoni.bank.curriculum import default_trees
    from dawoni.bank.scope import compose_scope

    trees = default_trees()
    a = _scope(
        "A교재 힘과 운동",
        ["phy1.1.1"],
        ["뉴턴 운동 법칙", "운동량"],
        skills=[Skill(id="S1", name="자유 물체도", description="-", concept_ids=["C1"])],
        forbidden=[ForbiddenTechnique(name="미분")],
        conventions=["g=10"],
    )
    b = _scope(
        "B교재 운동량",
        ["phy1.1.1.3"],
        ["운동량", "충격량"],
        skills=[
            Skill(id="S1", name="충격량-운동량 정리", description="-", concept_ids=["C2", "C1"])
        ],
        conventions=["g=10", "부호는 오른쪽 +"],
    )
    c = _scope("C교재 전자기", ["phy1.2.2"], ["전자기 유도"])
    res = compose_scope(
        [a, b, c], trees, ["phy1.1.1.3", "phy1.1.2.1"], title="운동량+역학적 에너지"
    )
    s = res.scope
    assert res.used == ["A교재 힘과 운동", "B교재 운동량"], "체크한 단원을 품은 교재 단원도 포함"
    assert [x.name for x in s.concepts] == [
        "뉴턴 운동 법칙",
        "운동량",
        "충격량",
        "역학적 에너지 보존",
    ]
    assert s.skills[1].concept_ids == ["C2", "C3"], "기법이 가리키는 개념 ID도 다시 매김"
    assert s.conventions == ["g=10", "부호는 오른쪽 +"] and [f.name for f in s.forbidden] == [
        "미분"
    ]
    assert any("역학적 에너지 보존" in w for w in res.warnings), (
        "교재 개념 없는 단원은 단원 이름으로"
    )
    assert s.unit_ids == ["phy1.1.1.3", "phy1.1.2.1"] and s.level == "고등학교 물리학 I"
    with pytest.raises(BankError, match="없는 단원"):
        compose_scope([a], trees, ["phy1.9"], title="x")


# ---------------------------------------------------------------- 문제 만들기 (문제은행 기준)


def test_variant_from_bank_item_goes_back_to_bank(
    app_paths: AppPaths, tmp_path: Path, make_fake: MakeFake
) -> None:
    from test_generate import M, T, solved
    from test_grading import FIX as GFIX
    from test_grading import RUBRIC_OK

    from dawoni.bank.make import approve_to_bank, generate_from_bank
    from dawoni.generate.workflow import GenStore

    tpl = {
        "item": {"item_type": "essay", "stem": [{"type": "paragraph", "runs": [
            T("상수 "), M(r"a \ge 0"), T("에 대하여 이차함수 "), M("y=x^2-2ax+7"), T("의 최솟값을 a로 나타내시오.")]}]},
        "answer_expr": "7 - a**2", "assumptions": ["a >= 0"], "note": "일차항 계수를 문자로",
    }  # fmt: skip
    script = [solved("3", "3"), tpl, solved("7-a^2", "7 - a**2"), RUBRIC_OK]
    exe = make_fake(
        "script", script=[{"response": json.dumps(r, ensure_ascii=False)} for r in script]
    )
    cfg = AppConfig()
    cfg.agents["agy"].executable = str(exe)
    cfg.agents["claude"].enabled = False
    cfg.agents["codex"].enabled = False
    save_config(app_paths, cfg)
    ws = create_workspace(tmp_path / "ws")
    ScopeStore(ws).save(ScopeProfile.model_validate(read_json(GFIX / "scope.json")))
    svc = open_services(ws.root, paths=app_paths)
    q = ExamQuestion.model_validate(read_json(GFIX / "questions.json")[0])
    base = BankItem(id="bk-base1", question=q, unit_ids=["phy1.1.1"],
                    source=BankSourceRef(title="OO 문제집", kind="problem_bank", pages=[5], number_label="7"),
                    answer_check=AnswerComparison(verdict="differs"))  # fmt: skip
    bank = BankStore(ws)
    bank.save(base)
    with pytest.raises(BankError, match="답 확인"):
        generate_from_bank(svc, "bk-base1", "symbolic", scope_id="scope-quadratic-minmax")
    ok = base.model_copy(
        update={"answer_check": AnswerComparison(verdict="differs", resolved="ai")}
    )
    bank.save(ok)

    sets = generate_from_bank(svc, "bk-base1", "symbolic", scope_id="scope-quadratic-minmax")
    assert len(fake_prompts(exe)) == 4, "고른 범위로 기준 문항을 먼저 풀고 제작"
    s = sets[0]
    assert s.source.bank_item_id == "bk-base1" and s.source.exam_id is None
    assert s.agent == "agy", "문항을 만든 AI 기록"
    assert GenStore(ws).for_bank_item("bk-base1") == [s]

    made = approve_to_bank(svc, s.id, by="김선생")
    assert made.origin == "variant" and made.parent_id == "bk-base1" and made.gen_kind == "symbolic"
    assert made.unit_ids == ["phy1.1.1"] and made.source.display().startswith(
        "다워니 제작 · 기준 OO 문제집"
    )
    assert made.answer_check and made.answer_check.verdict == "no_original" and not made.needs_check
    assert GenStore(ws).get(s.id).status == "approved"
    assert approve_to_bank(svc, s.id).id == made.id, "두 번 넣지 않음"


def test_variant_base_is_solved_first_when_unsolved(setup: Setup) -> None:
    from dawoni.bank.make import base_for

    svc, sid, exe = setup([_sol("③", ["C5"])])
    one, _ = bi.ingest_chunk(svc, sid, 1, machine="pc-a").items
    assert one.solution is None
    _, sol, warnings = base_for(svc, one, None)
    assert sol is not None and sol.final_answer.latex == "③" and warnings == []
    assert len(fake_prompts(exe)) == 1, "한 번 풀어 원본 답과 비교 (같으면 검산 없음)"
    saved = BankStore(svc.workspace).get(one.id)
    assert saved.solution is not None and saved.answer_check is not None, (
        "푼 결과는 문제은행에 저장"
    )
    assert saved.answer_check.verdict == "same"
